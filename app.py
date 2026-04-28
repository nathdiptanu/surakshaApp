from __future__ import annotations

import csv
import io
import json
import os
import shutil
import sqlite3
from datetime import datetime
from pathlib import Path
from typing import Iterable
from urllib.parse import quote
from uuid import uuid4

import qrcode
from flask import Flask, Response, jsonify, redirect, render_template, request, send_file, session, url_for
from openpyxl import load_workbook
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parent
DATA_DIR = ROOT / "data"
DB_PATH = DATA_DIR / "suretrace.db"
LEGACY_DB_PATH = DATA_DIR / ("sura" + "ksha.db")
APP_NAME = "SureTrace"
APP_CODE_PREFIX = "SRT"
CONFIG_PATH = ROOT / "instance" / "admin_credentials.json"

CATEGORIES = {
    "kids": {"label": "Kids", "accent": "#ff2f7d", "dark": "#4c0f7a", "light": "#fff04f"},
    "elderly": {"label": "Elderly", "accent": "#00a8ff", "dark": "#004aad", "light": "#93f5c4"},
    "bike": {"label": "Bike", "accent": "#ff7a00", "dark": "#1d1d1d", "light": "#ffe600"},
    "car": {"label": "Car", "accent": "#00c853", "dark": "#004d40", "light": "#b2ff59"},
    "employee": {"label": "Employee", "accent": "#7c4dff", "dark": "#25136d", "light": "#00e5ff"},
    "other": {"label": "Other", "accent": "#00b8d4", "dark": "#263238", "light": "#ffca28"},
}

FORMATS = {
    "pvc": {"label": "PVC Card", "width": 86, "height": 54},
    "jacket_tag": {"label": "Jacket Tag", "width": 70, "height": 42},
    "sticker_round": {"label": "Round Sticker", "width": 52, "height": 52},
    "vehicle_sticker": {"label": "Vehicle Sticker", "width": 92, "height": 38},
    "employee_badge": {"label": "Employee Badge", "width": 54, "height": 86},
}

FIELDS = [
    "category",
    "format",
    "name",
    "apartment",
    "location",
    "emergency_contact",
    "company",
    "notes",
]


app = Flask(__name__)
app.config["SEND_FILE_MAX_AGE_DEFAULT"] = 0


def load_private_config() -> dict:
    if CONFIG_PATH.exists():
        return json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    return {}


PRIVATE_CONFIG = load_private_config()
ADMIN_USERNAME = os.environ.get("SURETRACE_ADMIN_USERNAME") or PRIVATE_CONFIG.get("username")
ADMIN_PASSWORD = os.environ.get("SURETRACE_ADMIN_PASSWORD") or PRIVATE_CONFIG.get("password")
CLEANUP_PASSWORD = os.environ.get("SURETRACE_CLEANUP_PASSWORD") or PRIVATE_CONFIG.get("cleanup_password") or "bunty"
app.config["SECRET_KEY"] = (
    os.environ.get("SURETRACE_SECRET_KEY")
    or PRIVATE_CONFIG.get("secret_key")
    or "suretrace-dev-session-change-before-hosting"
)


@app.before_request
def require_login():
    public_paths = {"/login"}
    is_static = request.path.startswith("/static/")
    if request.path in public_paths or is_static:
        return None
    if not session.get("admin_logged_in"):
        if request.path.startswith("/api/"):
            return jsonify({"error": "Login required."}), 401
        return redirect(url_for("login", next=request.path))
    return None


def db() -> sqlite3.Connection:
    DATA_DIR.mkdir(exist_ok=True)
    if LEGACY_DB_PATH.exists() and not DB_PATH.exists():
        shutil.copy2(LEGACY_DB_PATH, DB_PATH)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db() -> None:
    with db() as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS qr_orders (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                code TEXT NOT NULL UNIQUE,
                category TEXT NOT NULL DEFAULT 'kids',
                name TEXT,
                apartment TEXT,
                location TEXT,
                emergency_contact TEXT,
                company TEXT,
                notes TEXT,
                created_at TEXT NOT NULL
            )
            """
        )
        columns = [row["name"] for row in conn.execute("PRAGMA table_info(qr_orders)").fetchall()]
        if "format" not in columns:
            conn.execute("ALTER TABLE qr_orders ADD COLUMN format TEXT NOT NULL DEFAULT 'pvc'")


def row_to_dict(row: sqlite3.Row) -> dict:
    return {key: row[key] for key in row.keys()}


def normalize_record(source: dict) -> dict:
    data = {field: (source.get(field) or "").strip() for field in FIELDS}
    data["category"] = data["category"].lower() if data["category"] else "kids"
    if data["category"] not in CATEGORIES:
        data["category"] = "other"
    data["format"] = data["format"].lower() if data["format"] else default_format(data["category"])
    if data["format"] not in FORMATS:
        data["format"] = default_format(data["category"])
    return data


def default_format(category: str) -> str:
    return {
        "kids": "jacket_tag",
        "elderly": "sticker_round",
        "bike": "vehicle_sticker",
        "car": "vehicle_sticker",
        "employee": "employee_badge",
    }.get(category, "pvc")


def validate_record(data: dict) -> list[str]:
    required = ["category", "format", "name", "apartment", "location", "emergency_contact", "notes"]
    labels = {
        "category": "Use case",
        "format": "Card format",
        "name": "Name",
        "apartment": "Apartment / address",
        "location": "Location",
        "emergency_contact": "Emergency contact",
        "notes": "Notes",
    }
    return [labels[field] for field in required if not data.get(field)]


def whatsapp_message(record: dict) -> str:
    lines = [
        "SureTrace QR scanned",
        f"ID: {record.get('code', '')}",
        f"Use case: {CATEGORIES.get(record.get('category'), CATEGORIES['other'])['label']}",
        f"Name: {record.get('name', '')}",
        f"Apartment / Address: {record.get('apartment', '')}",
        f"Location: {record.get('location', '')}",
        f"Emergency contact: {record.get('emergency_contact', '')}",
    ]
    if record.get("company"):
        lines.append(f"Company: {record.get('company')}")
    lines.append(f"Notes: {record.get('notes', '')}")
    return "\n".join(line for line in lines if not line.endswith(": "))


def public_payload(record: dict) -> str:
    return f"https://wa.me/?text={quote(whatsapp_message(record))}"


def make_qr_image(record: dict):
    qr = qrcode.QRCode(version=None, box_size=8, border=2)
    qr.add_data(public_payload(record))
    qr.make(fit=True)
    return qr.make_image(fill_color="#111111", back_color="white").convert("RGB")


def qr_png(record: dict) -> io.BytesIO:
    output = io.BytesIO()
    make_qr_image(record).save(output, format="PNG")
    output.seek(0)
    return output


def create_record(data: dict) -> dict:
    record = normalize_record(data)
    errors = validate_record(record)
    if errors:
        raise ValueError(", ".join(errors))
    record["code"] = f"{APP_CODE_PREFIX}-{datetime.now():%Y%m%d}-{uuid4().hex[:8].upper()}"
    record["created_at"] = datetime.now().isoformat(timespec="seconds")
    with db() as conn:
        cur = conn.execute(
            """
            INSERT INTO qr_orders
            (code, category, format, name, apartment, location, emergency_contact, company, notes, created_at)
            VALUES (:code, :category, :format, :name, :apartment, :location, :emergency_contact, :company, :notes, :created_at)
            """,
            record,
        )
        record["id"] = cur.lastrowid
    return record


def get_records(ids: Iterable[int] | None = None) -> list[dict]:
    with db() as conn:
        if ids:
            placeholders = ",".join("?" for _ in ids)
            rows = conn.execute(
                f"SELECT * FROM qr_orders WHERE id IN ({placeholders}) ORDER BY id DESC", list(ids)
            ).fetchall()
        else:
            rows = conn.execute("SELECT * FROM qr_orders ORDER BY id DESC").fetchall()
    return [row_to_dict(row) for row in rows]


def delete_all_records() -> int:
    with db() as conn:
        count = conn.execute("SELECT COUNT(*) FROM qr_orders").fetchone()[0]
        conn.execute("DELETE FROM qr_orders")
    return count


def draw_card(pdf: canvas.Canvas, record: dict, x: float, y: float) -> None:
    format_style = FORMATS.get(record.get("format"), FORMATS["pvc"])
    width = format_style["width"] * mm
    height = format_style["height"] * mm
    style = CATEGORIES.get(record["category"], CATEGORIES["other"])
    accent = colors.HexColor(style["accent"])
    dark = colors.HexColor(style["dark"])
    light = colors.HexColor(style["light"])
    is_round = record.get("format") == "sticker_round"

    pdf.setFillColor(light)
    if is_round:
        pdf.circle(x + width / 2, y + height / 2, width / 2, stroke=0, fill=1)
    else:
        pdf.roundRect(x, y, width, height, 8, stroke=0, fill=1)
    pdf.setFillColor(accent)
    pdf.roundRect(x, y + height - 14 * mm, width, 14 * mm, 8, stroke=0, fill=1)
    pdf.setFillColor(dark)
    pdf.rect(x, y + height - 7 * mm, width, 7 * mm, stroke=0, fill=1)

    pdf.setFillColor(colors.white)
    pdf.setFont("Helvetica-Bold", 13)
    draw_pdf_logo(pdf, x + 5 * mm, y + height - 10.5 * mm, 6 * mm)
    pdf.drawString(x + 13 * mm, y + height - 9 * mm, APP_NAME)
    pdf.setFont("Helvetica", 7)
    pdf.drawRightString(x + width - 5 * mm, y + height - 9 * mm, style["label"].upper())
    draw_pdf_category_icon(pdf, record["category"], x + width - 13 * mm, y + height - 12.5 * mm, 5 * mm)

    qr_img = make_qr_image(record)
    qr_size = min(width * 0.43, height * 0.52)
    if record.get("format") == "employee_badge":
        qr_size = 30 * mm
        qr_x = x + (width - qr_size) / 2
        qr_y = y + 25 * mm
        text_x = x + 5 * mm
        text_y = y + 18 * mm
    elif is_round:
        qr_size = 28 * mm
        qr_x = x + (width - qr_size) / 2
        qr_y = y + 12 * mm
        text_x = x + 7 * mm
        text_y = y + 9 * mm
    else:
        qr_x = x + 5 * mm
        qr_y = y + 7 * mm
        text_x = x + qr_size + 10 * mm
        text_y = y + height - 23 * mm
    pdf.drawInlineImage(qr_img, qr_x, qr_y, qr_size, qr_size)

    pdf.setFillColor(dark)
    pdf.setFont("Helvetica-Bold", 9)
    pdf.drawString(text_x, text_y, "SCAN FOR DETAILS")
    pdf.setFont("Helvetica", 6.7)
    pdf.drawString(text_x, text_y - 4 * mm, "Opens WhatsApp message")
    pdf.drawString(text_x, text_y - 8 * mm, f"ID: {record.get('code')}")

    pdf.setStrokeColor(accent)
    pdf.setLineWidth(1.6)
    if is_round:
        pdf.circle(x + width / 2, y + height / 2, width / 2, stroke=1, fill=0)
    else:
        pdf.roundRect(x, y, width, height, 8, stroke=1, fill=0)


def draw_pdf_logo(pdf: canvas.Canvas, x: float, y: float, size: float) -> None:
    pdf.saveState()
    pdf.setFillColor(colors.white)
    pdf.circle(x + size / 2, y + size / 2, size / 2, stroke=0, fill=1)
    pdf.setStrokeColor(colors.HexColor("#25136d"))
    pdf.setLineWidth(1.2)
    pdf.line(x + size * 0.28, y + size * 0.46, x + size * 0.47, y + size * 0.28)
    pdf.line(x + size * 0.47, y + size * 0.28, x + size * 0.74, y + size * 0.72)
    pdf.setFillColor(colors.HexColor("#ff2f7d"))
    pdf.circle(x + size * 0.28, y + size * 0.46, size * 0.09, stroke=0, fill=1)
    pdf.setFillColor(colors.HexColor("#00b8d4"))
    pdf.circle(x + size * 0.74, y + size * 0.72, size * 0.09, stroke=0, fill=1)
    pdf.restoreState()


def draw_pdf_category_icon(pdf: canvas.Canvas, category: str, x: float, y: float, size: float) -> None:
    pdf.saveState()
    pdf.setFillColor(colors.white)
    pdf.roundRect(x, y, size, size, 2, stroke=0, fill=1)
    pdf.setFillColor(colors.HexColor("#17202a"))
    if category == "car":
        pdf.roundRect(x + size * 0.18, y + size * 0.38, size * 0.64, size * 0.24, 2, stroke=0, fill=1)
        pdf.rect(x + size * 0.32, y + size * 0.58, size * 0.28, size * 0.14, stroke=0, fill=1)
        pdf.circle(x + size * 0.3, y + size * 0.32, size * 0.08, stroke=0, fill=1)
        pdf.circle(x + size * 0.7, y + size * 0.32, size * 0.08, stroke=0, fill=1)
    elif category == "bike":
        pdf.circle(x + size * 0.28, y + size * 0.34, size * 0.11, stroke=1, fill=0)
        pdf.circle(x + size * 0.72, y + size * 0.34, size * 0.11, stroke=1, fill=0)
        pdf.line(x + size * 0.28, y + size * 0.34, x + size * 0.5, y + size * 0.55)
        pdf.line(x + size * 0.5, y + size * 0.55, x + size * 0.72, y + size * 0.34)
    elif category == "kids":
        pdf.circle(x + size * 0.5, y + size * 0.62, size * 0.16, stroke=0, fill=1)
        pdf.roundRect(x + size * 0.28, y + size * 0.2, size * 0.44, size * 0.28, 3, stroke=0, fill=1)
    elif category == "elderly":
        pdf.circle(x + size * 0.42, y + size * 0.66, size * 0.12, stroke=0, fill=1)
        pdf.line(x + size * 0.42, y + size * 0.54, x + size * 0.42, y + size * 0.28)
        pdf.line(x + size * 0.55, y + size * 0.5, x + size * 0.74, y + size * 0.2)
    elif category == "employee":
        pdf.roundRect(x + size * 0.25, y + size * 0.18, size * 0.5, size * 0.62, 2, stroke=1, fill=0)
        pdf.circle(x + size * 0.5, y + size * 0.58, size * 0.09, stroke=0, fill=1)
        pdf.rect(x + size * 0.37, y + size * 0.32, size * 0.26, size * 0.08, stroke=0, fill=1)
    else:
        pdf.circle(x + size * 0.5, y + size * 0.5, size * 0.2, stroke=1, fill=0)
        pdf.line(x + size * 0.65, y + size * 0.35, x + size * 0.82, y + size * 0.18)
    pdf.restoreState()


def build_pdf(records: list[dict]) -> io.BytesIO:
    output = io.BytesIO()
    pdf = canvas.Canvas(output, pagesize=A4)
    page_w, page_h = A4
    margin_x = 13 * mm
    margin_y = 14 * mm
    gap_x = 10 * mm
    gap_y = 9 * mm
    card_w = 92 * mm
    card_h = 86 * mm
    x_positions = [margin_x, margin_x + card_w + gap_x]
    y = page_h - margin_y - card_h

    for idx, record in enumerate(records):
        col = idx % 2
        draw_card(pdf, record, x_positions[col], y)
        if col == 1:
            y -= card_h + gap_y
        if y < margin_y and idx != len(records) - 1:
            pdf.showPage()
            y = page_h - margin_y - card_h

    pdf.save()
    output.seek(0)
    return output


@app.route("/")
def index():
    return render_template("index.html", categories=CATEGORIES, formats=FORMATS)


@app.route("/login", methods=["GET", "POST"])
def login():
    error = ""
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()
        if not ADMIN_USERNAME or not ADMIN_PASSWORD:
            error = "Admin credentials are not configured on this server."
            return render_template("login.html", error=error), 500
        if username == ADMIN_USERNAME and password == ADMIN_PASSWORD:
            session["admin_logged_in"] = True
            return redirect(request.args.get("next") or url_for("index"))
        error = "Invalid username or password."
    return render_template("login.html", error=error)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/api/records", methods=["GET", "POST"])
def records_api():
    if request.method == "POST":
        try:
            record = create_record(request.get_json(force=True))
        except ValueError as exc:
            return jsonify({"error": f"Required fields missing: {exc}"}), 400
        return jsonify(record), 201
    return jsonify(get_records())


@app.route("/api/upload", methods=["POST"])
def upload_api():
    file = request.files.get("file")
    if not file:
        return jsonify({"error": "Upload an .xlsx or .csv file."}), 400

    created = []
    if file.filename.lower().endswith(".csv"):
        rows = csv.DictReader(io.StringIO(file.read().decode("utf-8-sig")))
        for row in rows:
            try:
                created.append(create_record({k.lower().strip().replace(" ", "_"): v for k, v in row.items()}))
            except ValueError:
                continue
    else:
        workbook = load_workbook(file, read_only=True, data_only=True)
        sheet = workbook.active
        headers = [str(cell.value or "").lower().strip().replace(" ", "_") for cell in next(sheet.iter_rows(max_row=1))]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            values = {headers[i]: str(value or "") for i, value in enumerate(row) if i < len(headers)}
            if any(values.values()):
                try:
                    created.append(create_record(values))
                except ValueError:
                    continue

    return jsonify({"created": created, "count": len(created)})


@app.route("/api/records/delete-all", methods=["POST"])
def delete_all_api():
    password = request.get_json(force=True).get("password", "")
    if password != CLEANUP_PASSWORD:
        return jsonify({"error": "Cleanup password is incorrect."}), 403
    deleted = delete_all_records()
    return jsonify({"deleted": deleted})


@app.route("/download/pdf")
def download_pdf():
    ids = [int(item) for item in request.args.get("ids", "").split(",") if item.isdigit()]
    records = get_records(ids or None)
    if not records:
        return jsonify({"error": "No records available for PDF."}), 404
    return send_file(
        build_pdf(records),
        mimetype="application/pdf",
        as_attachment=True,
        download_name=f"suretrace_qr_cards_{datetime.now():%Y%m%d_%H%M}.pdf",
    )


@app.route("/download/orders.csv")
def download_orders():
    records = get_records()
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=["id", "code", *FIELDS, "created_at"])
    writer.writeheader()
    writer.writerows(records)
    return Response(
        output.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=suretrace_orders.csv"},
    )


@app.route("/download/orders.xlsx")
def download_orders_xlsx():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SureTrace Orders"
    headers = ["id", "code", *FIELDS, "created_at"]
    sheet.append(headers)
    for record in get_records():
        sheet.append([record.get(header, "") for header in headers])
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 20

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="suretrace_orders.xlsx",
    )


@app.route("/qr/<int:record_id>.png")
def qr_image(record_id: int):
    records = get_records([record_id])
    if not records:
        return jsonify({"error": "QR record not found."}), 404
    return send_file(qr_png(records[0]), mimetype="image/png")


@app.route("/sample.csv")
def sample_csv():
    sample = io.StringIO()
    writer = csv.DictWriter(sample, fieldnames=FIELDS)
    writer.writeheader()
    writer.writerow(
        {
            "category": "kids",
            "format": "jacket_tag",
            "name": "Aarav Sharma",
            "apartment": "Green Heights A-304",
            "location": "Whitefield, Bengaluru",
            "emergency_contact": "+91 98765 43210",
            "company": "",
            "notes": "Blood group O+",
        }
    )
    writer.writerow(
        {
            "category": "employee",
            "format": "employee_badge",
            "name": "Employee Asset",
            "apartment": "",
            "location": "Pune Office",
            "emergency_contact": "",
            "company": "Acme Pvt Ltd",
            "notes": "Asset batch 001",
        }
    )
    return Response(
        sample.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=suretrace_sample_upload.csv"},
    )


init_db()


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=True)
